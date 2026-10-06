import csv
import os
import re
import unicodedata

from collections import Counter
from datetime import datetime, timezone

from bson import ObjectId
from bson.errors import InvalidId

from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from extensions import db


# ============================================================
# CONFIGURACIÓN
# ============================================================

MAX_EXTRACTED_CHARACTERS = 120000

SUMMARY_MAX_CHARACTERS = 650

MAX_KEYWORDS = 12


# ============================================================
# PALABRAS VACÍAS
# ============================================================

STOP_WORDS = {
    "a",
    "al",
    "algo",
    "algunas",
    "algunos",
    "ante",
    "antes",
    "como",
    "con",
    "contra",
    "cual",
    "cuando",
    "de",
    "del",
    "desde",
    "donde",
    "durante",
    "e",
    "el",
    "ella",
    "ellas",
    "ellos",
    "en",
    "entre",
    "era",
    "eran",
    "es",
    "esa",
    "esas",
    "ese",
    "eso",
    "esos",
    "esta",
    "estaba",
    "estaban",
    "este",
    "esto",
    "estos",
    "fue",
    "fueron",
    "ha",
    "han",
    "hasta",
    "hay",
    "la",
    "las",
    "lo",
    "los",
    "más",
    "mas",
    "me",
    "mi",
    "mis",
    "muy",
    "no",
    "nos",
    "o",
    "para",
    "pero",
    "por",
    "porque",
    "que",
    "qué",
    "se",
    "sin",
    "sobre",
    "son",
    "su",
    "sus",
    "también",
    "tambien",
    "te",
    "tiene",
    "tienen",
    "tu",
    "un",
    "una",
    "uno",
    "unos",
    "unas",
    "y",
    "ya"
}


# ============================================================
# EXCEPCIONES
# ============================================================

class DocumentContentError(Exception):
    pass


class UnsupportedDocumentError(
    DocumentContentError
):
    pass


class DocumentFileNotFoundError(
    DocumentContentError
):
    pass


# ============================================================
# NORMALIZAR TEXTO
# ============================================================

def normalize_text(value):

    text = str(
        value or ""
    )

    text = unicodedata.normalize(
        "NFD",
        text
    )

    text = "".join(
        character
        for character in text
        if unicodedata.category(
            character
        ) != "Mn"
    )

    text = text.lower()

    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text.strip()


# ============================================================
# LIMPIAR TEXTO EXTRAÍDO
# ============================================================

def clean_extracted_text(text):

    text = str(
        text or ""
    )

    text = text.replace(
        "\x00",
        " "
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    text = re.sub(
        r"\n\s*\n\s*\n+",
        "\n\n",
        text
    )

    text = text.strip()

    if len(text) > MAX_EXTRACTED_CHARACTERS:

        text = text[
            :MAX_EXTRACTED_CHARACTERS
        ]

    return text


# ============================================================
# LEER PDF
# ============================================================

def extract_pdf(
    full_path
):

    try:

        reader = PdfReader(
            full_path
        )

        parts = []

        for page in reader.pages:

            try:

                text = (
                    page.extract_text()
                    or ""
                )

            except Exception:

                text = ""

            if text.strip():

                parts.append(
                    text
                )

        return clean_extracted_text(
            "\n\n".join(
                parts
            )
        )

    except Exception as error:

        raise DocumentContentError(
            (
                "No fue posible leer el PDF: "
                f"{error}"
            )
        )


# ============================================================
# LEER DOCX
# ============================================================

def extract_docx(
    full_path
):

    try:

        document = WordDocument(
            full_path
        )

        parts = []

        # ----------------------------------------------------
        # PÁRRAFOS
        # ----------------------------------------------------

        for paragraph in document.paragraphs:

            text = (
                paragraph.text
                or ""
            ).strip()

            if text:

                parts.append(
                    text
                )

        # ----------------------------------------------------
        # TABLAS
        # ----------------------------------------------------

        for table in document.tables:

            for row in table.rows:

                values = []

                for cell in row.cells:

                    value = (
                        cell.text
                        or ""
                    ).strip()

                    if value:

                        values.append(
                            value
                        )

                if values:

                    parts.append(
                        " | ".join(
                            values
                        )
                    )

        return clean_extracted_text(
            "\n".join(
                parts
            )
        )

    except Exception as error:

        raise DocumentContentError(
            (
                "No fue posible leer el documento "
                f"Word: {error}"
            )
        )


# ============================================================
# LEER TXT
# ============================================================

def extract_txt(
    full_path
):

    encodings = [
        "utf-8",
        "utf-8-sig",
        "latin-1",
        "cp1252"
    ]

    for encoding in encodings:

        try:

            with open(
                full_path,
                "r",
                encoding=encoding
            ) as file:

                return clean_extracted_text(
                    file.read()
                )

        except UnicodeDecodeError:

            continue

        except Exception as error:

            raise DocumentContentError(
                (
                    "No fue posible leer el archivo "
                    f"de texto: {error}"
                )
            )

    raise DocumentContentError(
        "No fue posible determinar la codificación del archivo."
    )


# ============================================================
# LEER CSV
# ============================================================

def extract_csv(
    full_path
):

    encodings = [
        "utf-8",
        "utf-8-sig",
        "latin-1",
        "cp1252"
    ]

    for encoding in encodings:

        try:

            parts = []

            with open(
                full_path,
                "r",
                encoding=encoding,
                newline=""
            ) as file:

                reader = csv.reader(
                    file
                )

                for row in reader:

                    values = [
                        str(value).strip()
                        for value in row
                        if str(value).strip()
                    ]

                    if values:

                        parts.append(
                            " | ".join(
                                values
                            )
                        )

            return clean_extracted_text(
                "\n".join(
                    parts
                )
            )

        except UnicodeDecodeError:

            continue

        except Exception as error:

            raise DocumentContentError(
                (
                    "No fue posible leer el CSV: "
                    f"{error}"
                )
            )

    raise DocumentContentError(
        "No fue posible determinar la codificación del CSV."
    )


# ============================================================
# LEER XLSX
# ============================================================

def extract_xlsx(
    full_path
):

    try:

        workbook = load_workbook(
            full_path,
            read_only=True,
            data_only=True
        )

        parts = []

        for sheet in workbook.worksheets:

            parts.append(
                f"Hoja: {sheet.title}"
            )

            for row in sheet.iter_rows(
                values_only=True
            ):

                values = [
                    str(value).strip()
                    for value in row
                    if value is not None
                    and str(value).strip()
                ]

                if values:

                    parts.append(
                        " | ".join(
                            values
                        )
                    )

                if (
                    sum(
                        len(part)
                        for part in parts
                    )
                    >= MAX_EXTRACTED_CHARACTERS
                ):

                    break

            if (
                sum(
                    len(part)
                    for part in parts
                )
                >= MAX_EXTRACTED_CHARACTERS
            ):

                break

        workbook.close()

        return clean_extracted_text(
            "\n".join(
                parts
            )
        )

    except Exception as error:

        raise DocumentContentError(
            (
                "No fue posible leer el Excel: "
                f"{error}"
            )
        )


# ============================================================
# EXTRAER TEXTO SEGÚN EXTENSIÓN
# ============================================================

def extract_text_from_file(
    full_path,
    file_type
):

    if not full_path:

        raise DocumentFileNotFoundError(
            "No se recibió una ruta de archivo."
        )

    if not os.path.isfile(
        full_path
    ):

        raise DocumentFileNotFoundError(
            (
                "No se encontró el archivo físico: "
                f"{full_path}"
            )
        )

    extension = str(
        file_type or ""
    ).lower().strip(".")

    if extension == "pdf":

        return extract_pdf(
            full_path
        )

    if extension == "docx":

        return extract_docx(
            full_path
        )

    if extension == "txt":

        return extract_txt(
            full_path
        )

    if extension == "csv":

        return extract_csv(
            full_path
        )

    if extension == "xlsx":

        return extract_xlsx(
            full_path
        )

    raise UnsupportedDocumentError(
        (
            "El análisis de contenido todavía "
            f"no está disponible para .{extension}"
        )
    )


# ============================================================
# DIVIDIR EN ORACIONES
# ============================================================

def split_sentences(
    text
):

    text = clean_extracted_text(
        text
    )

    if not text:

        return []

    sentences = re.split(
        r"(?<=[.!?])\s+|\n+",
        text
    )

    return [
        sentence.strip()
        for sentence in sentences
        if len(
            sentence.strip()
        ) >= 20
    ]


# ============================================================
# GENERAR RESUMEN
# ============================================================

def generate_summary(
    text,
    max_characters=SUMMARY_MAX_CHARACTERS
):

    text = clean_extracted_text(
        text
    )

    if not text:

        return (
            "No fue posible obtener suficiente texto "
            "para generar un resumen."
        )

    sentences = split_sentences(
        text
    )

    if not sentences:

        return text[
            :max_characters
        ]

    selected = []

    total_characters = 0

    # --------------------------------------------------------
    # PARA ESTA PRIMERA VERSIÓN TOMAMOS LAS PRIMERAS
    # ORACIONES ÚTILES DEL DOCUMENTO.
    # --------------------------------------------------------

    for sentence in sentences:

        if total_characters + len(
            sentence
        ) > max_characters:

            break

        selected.append(
            sentence
        )

        total_characters += (
            len(sentence)
            +
            1
        )

        if len(selected) >= 5:

            break

    summary = " ".join(
        selected
    ).strip()

    if not summary:

        summary = text[
            :max_characters
        ]

    if (
        len(text) > len(summary)
        and
        not summary.endswith(
            ("...", ".", "!", "?")
        )
    ):

        summary += "..."

    return summary


# ============================================================
# EXTRAER PALABRAS
# ============================================================

def tokenize(
    text
):

    normalized = normalize_text(
        text
    )

    words = re.findall(
        r"\b[a-z0-9][a-z0-9_-]{2,}\b",
        normalized
    )

    result = []

    for word in words:

        if word in STOP_WORDS:
            continue

        if word.isdigit():
            continue

        if len(word) < 3:
            continue

        result.append(
            word
        )

    return result


# ============================================================
# GENERAR PALABRAS CLAVE
# ============================================================

def generate_keywords(
    text,
    limit=MAX_KEYWORDS
):

    words = tokenize(
        text
    )

    if not words:

        return []

    counter = Counter(
        words
    )

    keywords = [
        word
        for word, _ in counter.most_common(
            limit
        )
    ]

    return keywords


# ============================================================
# RESOLVER ARCHIVO FÍSICO
# ============================================================

def resolve_document_path(
    document,
    root_path,
    upload_folder
):

    upload_dir = os.path.join(
        root_path,
        upload_folder
    )

    candidates = []

    stored_filename = document.get(
        "stored_filename"
    )

    file_path = document.get(
        "file_path"
    )

    if stored_filename:

        candidates.append(
            os.path.join(
                upload_dir,
                os.path.basename(
                    stored_filename
                )
            )
        )

    if file_path:

        candidates.append(
            os.path.join(
                upload_dir,
                os.path.basename(
                    file_path
                )
            )
        )

        candidates.append(
            os.path.join(
                root_path,
                file_path
            )
        )

        candidates.append(
            os.path.join(
                root_path,
                "static",
                file_path
            )
        )

    for candidate in candidates:

        normalized = os.path.abspath(
            candidate
        )

        if os.path.isfile(
            normalized
        ):

            return normalized

    raise DocumentFileNotFoundError(
        (
            "No fue posible localizar físicamente "
            "el archivo del documento."
        )
    )


# ============================================================
# ANALIZAR DOCUMENTO
# ============================================================

def analyze_document(
    document,
    root_path,
    upload_folder
):

    if not document:

        raise DocumentContentError(
            "Documento inválido."
        )

    file_type = str(
        document.get(
            "file_type",
            ""
        )
    ).lower().strip(".")

    supported = {
        "pdf",
        "docx",
        "txt",
        "csv",
        "xlsx"
    }

    if file_type not in supported:

        return {
            "status":
                "unsupported",

            "extracted_text":
                "",

            "summary":
                (
                    "El análisis de contenido no está "
                    f"disponible para archivos .{file_type}."
                ),

            "keywords":
                []
        }

    full_path = resolve_document_path(
        document=document,
        root_path=root_path,
        upload_folder=upload_folder
    )

    extracted_text = extract_text_from_file(
        full_path=full_path,
        file_type=file_type
    )

    if not extracted_text.strip():

        return {
            "status":
                "empty",

            "extracted_text":
                "",

            "summary":
                (
                    "El archivo pudo abrirse, pero no se "
                    "encontró texto extraíble. Si se trata "
                    "de un PDF escaneado será necesario OCR."
                ),

            "keywords":
                []
        }

    summary = generate_summary(
        extracted_text
    )

    keywords = generate_keywords(
        extracted_text
    )

    return {
        "status":
            "indexed",

        "extracted_text":
            extracted_text,

        "summary":
            summary,

        "keywords":
            keywords
    }


# ============================================================
# VALIDAR ID
# ============================================================

def validate_document_id(
    document_id
):

    try:

        return ObjectId(
            str(document_id)
        )

    except (
        InvalidId,
        TypeError,
        ValueError
    ):

        raise DocumentContentError(
            "El identificador del documento no es válido."
        )


# ============================================================
# INDEXAR DOCUMENTO Y GUARDAR EN MONGODB
# ============================================================

def index_document(
    document_id,
    root_path,
    upload_folder
):

    document_oid = validate_document_id(
        document_id
    )

    document = db.documents.find_one(
        {
            "_id":
                document_oid
        }
    )

    if not document:

        raise DocumentContentError(
            "Documento no encontrado."
        )

    if document.get(
        "is_deleted",
        False
    ):

        raise DocumentContentError(
            (
                "No se indexan documentos "
                "que están en la papelera."
            )
        )

    try:

        analysis = analyze_document(
            document=document,
            root_path=root_path,
            upload_folder=upload_folder
        )

        now = datetime.now(
            timezone.utc
        )

        update_data = {

            "extracted_text":
                analysis[
                    "extracted_text"
                ],

            "ai_summary":
                analysis[
                    "summary"
                ],

            "ai_keywords":
                analysis[
                    "keywords"
                ],

            "ai_index_status":
                analysis[
                    "status"
                ],

            "ai_indexed_at":
                now

        }

        db.documents.update_one(
            {
                "_id":
                    document_oid
            },
            {
                "$set":
                    update_data
            }
        )

        return {
            **analysis,
            "document_id":
                str(
                    document_oid
                )
        }

    except UnsupportedDocumentError as error:

        db.documents.update_one(
            {
                "_id":
                    document_oid
            },
            {
                "$set": {

                    "ai_index_status":
                        "unsupported",

                    "ai_index_error":
                        str(error),

                    "ai_indexed_at":
                        datetime.now(
                            timezone.utc
                        )

                }
            }
        )

        raise

    except Exception as error:

        db.documents.update_one(
            {
                "_id":
                    document_oid
            },
            {
                "$set": {

                    "ai_index_status":
                        "error",

                    "ai_index_error":
                        str(error),

                    "ai_indexed_at":
                        datetime.now(
                            timezone.utc
                        )

                }
            }
        )

        raise


# ============================================================
# REINDEXAR TODOS LOS DOCUMENTOS
# ============================================================

def index_all_documents(
    root_path,
    upload_folder
):

    documents = list(
        db.documents.find(
            {
                "is_deleted": {
                    "$ne":
                        True
                }
            }
        )
    )

    result = {

        "total":
            len(
                documents
            ),

        "indexed":
            0,

        "empty":
            0,

        "unsupported":
            0,

        "errors":
            0,

        "details":
            []

    }

    for document in documents:

        document_id = document.get(
            "_id"
        )

        title = document.get(
            "title",
            "Sin título"
        )

        try:

            analysis = index_document(
                document_id=document_id,
                root_path=root_path,
                upload_folder=upload_folder
            )

            status = analysis.get(
                "status"
            )

            if status == "indexed":

                result[
                    "indexed"
                ] += 1

            elif status == "empty":

                result[
                    "empty"
                ] += 1

            else:

                result[
                    "unsupported"
                ] += 1

            result[
                "details"
            ].append(
                {
                    "id":
                        str(
                            document_id
                        ),

                    "title":
                        title,

                    "status":
                        status
                }
            )

        except UnsupportedDocumentError:

            result[
                "unsupported"
            ] += 1

            result[
                "details"
            ].append(
                {
                    "id":
                        str(
                            document_id
                        ),

                    "title":
                        title,

                    "status":
                        "unsupported"
                }
            )

        except Exception as error:

            result[
                "errors"
            ] += 1

            result[
                "details"
            ].append(
                {
                    "id":
                        str(
                            document_id
                        ),

                    "title":
                        title,

                    "status":
                        "error",

                    "error":
                        str(
                            error
                        )
                }
            )

    return result